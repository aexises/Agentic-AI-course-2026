from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / "STUDYBOOK.md"
OUTPUT_DIR = ROOT / "output" / "studybook"
OUTPUT = OUTPUT_DIR / "agentic-ai-studybook.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
MUTED = "5B6573"
LIGHT = "E8EEF5"
GRID = "C8D1DC"


def set_font(run, name: str, size: float | None = None, bold: bool | None = None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    total = sum(widths)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_layout = tbl_pr.first_child_found_in("w:tblLayout")
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")

    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            tc_w.set(qn("w:w"), str(widths[index]))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(widths[index] / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def add_numbering(document, num_format: str, text: str, num_id: int, abstract_id: int):
    numbering = document.part.numbering_part.element
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)

    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    fmt = OxmlElement("w:numFmt")
    fmt.set(qn("w:val"), num_format)
    level.append(fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), text)
    level.append(lvl_text)
    justify = OxmlElement("w:lvlJc")
    justify.set(qn("w:val"), "left")
    level.append(justify)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "270")
    p_pr.append(ind)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "300")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.append(spacing)
    level.append(p_pr)
    abstract.append(level)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)


def apply_num(paragraph, num_id: int):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num)
    p_pr.append(num_pr)


def clean_inline(text: str) -> str:
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    text = text.replace("`", "")
    text = text.replace("**", "")
    return text.strip()


def configure_document(document: Document):
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        style = document.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header
    hp = header.paragraphs[0]
    hp.text = "AGENTIC AI  /  STUDYBOOK"
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp_run = hp.runs[0]
    set_font(hp_run, "Calibri", 8.5, True)
    hp_run.font.color.rgb = RGBColor.from_string(MUTED)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = fp.add_run("PAGE ")
    set_font(run, "Calibri", 8.5, True)
    run.font.color.rgb = RGBColor.from_string(MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    fp._p.append(fld)


def add_cover(document: Document):
    p = document.add_paragraph()
    p.paragraph_format.space_before = Pt(96)
    p.paragraph_format.space_after = Pt(20)
    run = p.add_run("AGENTIC AI")
    set_font(run, "Calibri", 34, True)
    run.font.color.rgb = RGBColor.from_string(BLUE)

    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(16)
    run = p.add_run("Studybook and Exam Conspect")
    set_font(run, "Calibri", 24, True)
    run.font.color.rgb = RGBColor.from_string(DARK_BLUE)

    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(32)
    run = p.add_run("From language models to reliable autonomous systems")
    set_font(run, "Calibri", 15, False)
    run.font.color.rgb = RGBColor.from_string(MUTED)

    table = document.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHT)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(
        "A self-contained learning edition derived only from the supplied presentations "
        "and their embedded references. Lecturer and institution identifiers were removed."
    )
    set_font(run, "Calibri", 11, False)
    document.add_page_break()


def add_markdown_table(document: Document, rows: list[list[str]]):
    if len(rows) < 2:
        return
    column_count = max(len(row) for row in rows)
    normalized = [row + [""] * (column_count - len(row)) for row in rows]
    table = document.add_table(rows=len(normalized), cols=column_count)
    widths = [9360 // column_count] * column_count
    widths[-1] += 9360 - sum(widths)
    for r_index, row in enumerate(normalized):
        for c_index, value in enumerate(row):
            cell = table.cell(r_index, c_index)
            cell.text = clean_inline(value)
            if r_index == 0:
                set_cell_shading(cell, LIGHT)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    set_font(run, "Calibri", 9.5, r_index == 0)
    set_table_geometry(table, widths)
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def build():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    document = Document()
    configure_document(document)
    add_numbering(document, "bullet", "•", num_id=91, abstract_id=91)
    add_numbering(document, "decimal", "%1.", num_id=92, abstract_id=92)
    add_cover(document)

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    index = 0
    skipped_document_title = False
    while index < len(lines):
        raw = lines[index].rstrip()
        line = raw.strip()
        if not line:
            index += 1
            continue
        if line == "---":
            index += 1
            continue
        if line.startswith("|"):
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            parsed = [
                [cell.strip() for cell in row.strip("|").split("|")]
                for row in table_lines
                if not re.match(r"^\|?\s*:?-+", row)
            ]
            add_markdown_table(document, parsed)
            continue
        if line.startswith("# "):
            if not skipped_document_title:
                skipped_document_title = True
                index += 1
                continue
            paragraph = document.add_paragraph(style="Heading 1")
            paragraph.paragraph_format.page_break_before = True
            paragraph.add_run(clean_inline(line[2:]))
            index += 1
            continue
        if line.startswith("## "):
            document.add_paragraph(clean_inline(line[3:]), style="Heading 2")
            index += 1
            continue
        if line.startswith("### "):
            document.add_paragraph(clean_inline(line[4:]), style="Heading 3")
            index += 1
            continue
        if line.startswith("> "):
            table = document.add_table(rows=1, cols=1)
            set_table_geometry(table, [9360])
            cell = table.cell(0, 0)
            set_cell_shading(cell, LIGHT)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(clean_inline(line[2:]))
            set_font(run, "Calibri", 11, True)
            index += 1
            continue
        if line.startswith("- "):
            paragraph = document.add_paragraph()
            apply_num(paragraph, 91)
            paragraph.add_run(clean_inline(line[2:]))
            index += 1
            continue
        number_match = re.match(r"^\d+\.\s+(.*)$", line)
        if number_match:
            paragraph = document.add_paragraph()
            apply_num(paragraph, 92)
            paragraph.add_run(clean_inline(number_match.group(1)))
            index += 1
            continue
        paragraph = document.add_paragraph()
        paragraph.add_run(clean_inline(line))
        index += 1

    core = document.core_properties
    core.title = "Agentic AI Studybook and Exam Conspect"
    core.subject = "Agentic AI course notes and exam preparation"
    core.author = ""
    core.last_modified_by = ""
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()

